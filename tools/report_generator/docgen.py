# -*- coding: utf-8 -*-
"""Генерация отчётов по практическим работам (python-docx + пост-обработка в MS Word).

Контент описывается списком блоков:
  ("h1", text)                      — заголовок раздела
  ("h2", text)                      — подзаголовок
  ("p", text)                       — абзац (поддерживается **жирный** фрагмент)
  ("pn", text)                      — абзац без красной строки
  ("ul", [items])                   — маркированный список (тире)
  ("ol", [items])                   — нумерованный список 1), 2) ...
  ("table", num, caption, header, rows, widths_cm)
        row может быть списком строк ячеек или ("group", text) — строка-раздел на всю ширину
  ("fig", path, caption, width_cm)
  ("pb",)                           — разрыв страницы
"""
import os, re, copy, json, time
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
TEXT_WIDTH_CM = 16.5


# ---------------------------------------------------------------- стили
def _set_font(run_or_style, size=14, bold=None, italic=None):
    f = run_or_style.font
    f.name = FONT
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    el = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.different_first_page_header_footer = True

    st = doc.styles["Normal"]
    _set_font(st, 14)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # язык — русский (переносы/проверка орфографии)
    rpr = st.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    rpr.append(lang)

    for name, size in (("Heading 1", 14), ("Heading 2", 14)):
        hs = doc.styles[name]
        _set_font(hs, size, bold=True, italic=False)
        hs.font.color.rgb = None
        hpf = hs.paragraph_format
        hpf.space_before = Pt(12 if name == "Heading 1" else 6)
        hpf.space_after = Pt(6)
        hpf.keep_with_next = True
        hpf.line_spacing = 1.5
        hpf.first_line_indent = Cm(1.25)
        hpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # убрать цвет темы
        r = hs.element.get_or_add_rPr()
        for c in r.findall(qn("w:color")):
            r.remove(c)

    # номер страницы в нижнем колонтитуле (кроме титульного листа)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Cm(0)
    _add_field(fp, "PAGE")
    return doc


def _add_field(paragraph, code):
    run = paragraph.add_run()
    _set_font(run, 12)
    for tag, text in (("begin", None), (None, code), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)


# ---------------------------------------------------------------- титульный лист
def title_page(doc, cfg, work_no, work_title, topic=None):
    def line(text="", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            _add_rich(p, text, size=size, bold=bold)
        return p

    line("Министерство науки и высшего образования РФ")
    line("Федеральное государственное бюджетное образовательное учреждение высшего образования")
    line("«Тульский государственный университет»")
    line("Институт прикладной математики и компьютерных наук", space_after=6)
    if cfg.get("department"):
        line(cfg["department"])
    for _ in range(5):
        line()
    for t in work_title.upper().split("\n"):
        line(t, bold=True)
    line()
    line(f"Отчет по выполнению практической работы № {work_no}")
    line("по курсу «**Программная инженерия**»")
    if topic:
        line()
        line(f"Тема: «{topic}»")
    for _ in range(4 if topic else 6):
        line()
    tab_p = lambda left, right: _two_side_line(doc, left, right)
    tab_p(f"Выполнил ст. гр. {cfg['group']}", cfg["student"])
    line()
    tab_p("Проверил преп.", cfg["teacher"])
    n_empty = 7 if topic else 9
    for _ in range(n_empty):
        line()
    p = line(f"Тула {cfg['year']}")
    p.paragraph_format.keep_with_next = False
    # разрыв страницы
    p.add_run().add_break(WD_BREAK.PAGE)


def _two_side_line(doc, left, right):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # правый табулятор на краю текста
    tabs = p.paragraph_format.tab_stops
    from docx.enum.text import WD_TAB_ALIGNMENT
    tabs.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    _add_rich(p, left + "\t" + right)
    return p


# ---------------------------------------------------------------- текст
_BOLD_RE = re.compile(r"(\*\*.+?\*\*|__.+?__)")


def _add_rich(p, text, size=14, bold=False, italic=False):
    for part in _BOLD_RE.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("__") and part.endswith("__"):
            part, i = part[2:-2], True
        r = p.add_run(part)
        _set_font(r, size, bold=b, italic=i)


def para(doc, text, indent=True, align=None, keep_next=False, size=14):
    p = doc.add_paragraph()
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if align is not None:
        p.alignment = align
    if keep_next:
        p.paragraph_format.keep_with_next = True
    _add_rich(p, text, size=size)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    _add_rich(p, text, bold=True)
    return p


def bullet_list(doc, items, numbered=False):
    for i, it in enumerate(items, 1):
        mark = f"{i}) " if numbered else "– "
        p = doc.add_paragraph()
        _add_rich(p, mark + it)
        # последний пункт можно не держать со следующим; первые — держим вместе с предыдущим абзацем


# ---------------------------------------------------------------- таблицы
def _cell_text(cell, text, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=False, bold=False):
    cell.text = ""
    lines = text.split("\n") if text else [""]
    for k, ln in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.widow_control = False
        p.alignment = align
        if keep_next:
            pf.keep_with_next = True
        _add_rich(p, ln, size=size, bold=bold)


def _row_cant_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trPr.append(th)


def _set_tbl_caption(table, value):
    tblPr = table._tbl.tblPr
    cap = OxmlElement("w:tblCaption")
    cap.set(qn("w:val"), value)
    tblPr.append(cap)


def _set_col_widths(table, widths_cm):
    total = sum(widths_cm)
    widths = [w * TEXT_WIDTH_CM / total for w in widths_cm]
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    return widths


def add_table_piece(doc, header, rows, widths_cm, tag, continuation_of=None):
    """Одна «часть» таблицы. continuation_of — номер таблицы для надписи «Продолжение таблицы»."""
    if continuation_of:
        p = para(doc, f"Продолжение таблицы {continuation_of}", indent=False,
                 align=WD_ALIGN_PARAGRAPH.RIGHT, keep_next=True)
    ncols = len(header)
    t = doc.add_table(rows=1 + len(rows), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_col_widths(t, widths_cm)
    hdr = t.rows[0]
    _row_cant_split(hdr, header=True)
    for c, txt in zip(hdr.cells, header):
        _cell_text(c, txt, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    for ri, r in enumerate(rows, 1):
        row = t.rows[ri]
        _row_cant_split(row)
        if isinstance(r, tuple) and r[0] == "group":
            m = row.cells[0].merge(row.cells[-1])
            _cell_text(m, r[1], align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        else:
            for c, txt in zip(row.cells, r):
                _cell_text(c, txt)
    _set_tbl_caption(t, tag)
    return t


# ---------------------------------------------------------------- сборка
def build(blocks, out_path, cfg, work_no, work_title, topic=None, splits=None):
    """splits: {table_num: [индексы строк данных, перед которыми начинается новая часть]}"""
    splits = splits or {}
    doc = new_document()
    title_page(doc, cfg, work_no, work_title, topic)
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            heading(doc, b[1], 1)
        elif kind == "h2":
            heading(doc, b[1], 2)
        elif kind == "p":
            para(doc, b[1])
        elif kind == "pk":  # абзац, связанный со следующим
            para(doc, b[1], keep_next=True)
        elif kind == "pn":
            para(doc, b[1], indent=False)
        elif kind == "quote":  # дословная выдержка из нормативного акта
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            _add_rich(p, "«" + b[1] + "»", italic=True)
            if len(b) > 2:
                _add_rich(p, " (" + b[2] + ")")
        elif kind == "ul":
            bullet_list(doc, b[1])
        elif kind == "ol":
            bullet_list(doc, b[1], numbered=True)
        elif kind == "src":  # список источников — без растягивания по ширине (длинные URL)
            for i, it in enumerate(b[1], 1):
                para(doc, f"{i}. {it}", align=WD_ALIGN_PARAGRAPH.LEFT)
        elif kind == "table":
            _, num, caption, header, rows, widths = b
            para(doc, f"Таблица {num} – {caption}", indent=False,
                 align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
            cuts = [0] + sorted(splits.get(num, [])) + [len(rows)]
            for k in range(len(cuts) - 1):
                part = rows[cuts[k]:cuts[k + 1]]
                add_table_piece(doc, header, part, widths, f"{num}|{cuts[k]}",
                                continuation_of=num if k else None)
            # небольшой отступ после таблицы
            sp = doc.add_paragraph()
            sp.paragraph_format.line_spacing = 1.0
        elif kind == "fig":
            _, path, caption, width = b
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(6)
            w, h = _img_size(path)
            width = min(width, TEXT_WIDTH_CM)
            max_h = 21.0
            if h / w * width > max_h:
                width = max_h * w / h
            p.add_run().add_picture(path, width=Cm(width))
            c = para(doc, caption, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
            c.paragraph_format.space_after = Pt(6)
        elif kind == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError(kind)
    doc.save(out_path)


def _img_size(path):
    from PIL import Image
    with Image.open(path) as im:
        return im.size


# ---------------------------------------------------------------- MS Word: разбивка таблиц по страницам
def word_find_splits(word, path):
    """Анализ разбиения таблиц по страницам.
    Возвращает список изменений [(порядок таблицы, номер, 'add'|'del', индекс строки)] в порядке документа."""
    doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True)
    try:
        doc.Repaginate()
        changes = []
        prev = {}  # num -> страница последней строки предыдущей части
        for i in range(1, doc.Tables.Count + 1):
            t = doc.Tables(i)
            title = t.Title or ""
            if "|" not in title:
                continue
            num, start = title.split("|")
            start = int(start)
            n = t.Rows.Count
            first_page = t.Rows(1).Range.Information(3)
            # лишний разрез: продолжение начинается на той же странице, где закончилась предыдущая часть
            if start > 0 and prev.get(num) == first_page:
                changes.append((i, num, "del", start))
            for r in range(2, n + 1):
                pg = t.Rows(r).Range.Information(3)
                if pg != first_page:
                    if r > 2:
                        changes.append((i, num, "add", start + (r - 2)))
                    break
            prev[num] = t.Rows(n).Range.Information(3)
        return changes
    finally:
        doc.Close(False)


def word_export_pdf(word, path, pdf_path):
    doc = word.Documents.Open(os.path.abspath(path))
    try:
        doc.Repaginate()
        doc.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)
        pages = doc.ComputeStatistics(2)
    finally:
        doc.Close(False)
    return pages


def build_final(blocks, out_path, cfg, work_no, work_title, topic=None, word=None, pdf=True):
    """Итеративная сборка: Word определяет, где таблицы переходят на новую страницу,
    и таблица разрезается с надписью «Продолжение таблицы N» и повтором шапки."""
    import win32com.client
    own = word is None
    if own:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
    splits = {}
    try:
        for it in range(40):
            build(blocks, out_path, cfg, work_no, work_title, topic, splits)
            changes = word_find_splits(word, out_path)
            if not changes:
                break
            # обрабатываем только первое изменение по порядку документа — остальные пересчитаются
            _, num, kind, idx = changes[0]
            lst = splits.setdefault(num, [])
            if kind == "add" and idx not in lst and idx > 0:
                lst.append(idx)
            elif kind == "del" and idx in lst:
                lst.remove(idx)
            else:
                break
        pages = None
        if pdf:
            pages = word_export_pdf(word, out_path, os.path.splitext(out_path)[0] + ".pdf")
        return splits, pages
    finally:
        if own:
            word.Quit()
